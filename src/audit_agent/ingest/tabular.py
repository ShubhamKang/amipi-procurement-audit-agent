"""XLSX / CSV loaders with sheet+row provenance and per-row failure reporting."""
from __future__ import annotations

import csv
from pathlib import Path
from typing import Any, Callable, TypeVar

from openpyxl import load_workbook

from ..models import (
    Evidence,
    HistoryRecord,
    IngestionIssue,
    POLine,
    QuoteRecord,
    RateCard,
    Severity,
)
from .parsing import is_blank, norm_header, to_date, to_decimal, to_text

T = TypeVar("T")

# canonical field -> accepted header aliases (normalized)
PO_COLUMNS = {
    "po_id": ["poid", "po", "ponumber", "purchaseorder"],
    "line_id": ["lineid", "line", "poline"],
    "product_id": ["productid", "sku", "itemid", "product"],
    "description": ["description", "desc", "item"],
    "category": ["category", "type"],
    "qty": ["qty", "quantity"],
    "metal_g": ["metalg", "metalgrams", "grams", "metalweightg"],
    "stone_ct": ["stonect", "carat", "stonecarat"],
    "set_pcs": ["setpcs", "settings", "settingpcs", "stones"],
    "finish": ["finish"],
    "due_date": ["duedate", "due", "deliverydate"],
}
RATE_COLUMNS = {
    "rate_id": ["rateid", "id"],
    "vendor": ["vendor", "supplier"],
    "effective_from": ["effectivefrom", "from", "validfrom", "startdate"],
    "effective_to": ["effectiveto", "to", "validto", "enddate"],
    "category": ["category"],
    "metal": ["metal", "metalgrade", "grade"],
    "metal_usd_per_gram": ["metalusdpergram", "metalrate"],
    "setting_usd_per_stone": ["settingusdperstone", "settingrate"],
    "labor_usd_per_unit": ["laborusdperunit", "labourusdperunit", "laborrate"],
    "handling_rate": ["handlingrate", "handling", "handlingfraction"],
    "finish_rhodium_usd_per_unit": ["finishrhodiumusdperunit", "rhodiumusdperunit", "rhodium"],
    "notes": ["notes", "note", "comment"],
}
QUOTE_COLUMNS = {
    "quote_id": ["quoteid", "quote", "id"],
    "vendor": ["vendor", "supplier"],
    "quote_date": ["quotedate", "date", "quoteddate"],
    "po_id": ["poid", "po"],
    "supplier_line": ["supplierline", "line"],
    "vendor_sku": ["vendorsku", "sku", "productid"],
    "description": ["description", "item"],
    "qty": ["qty", "quantity"],
    "unit": ["unit", "uom"],
    "quoted_total_usd": ["quotedtotalusd", "total", "quotedtotal", "amount"],
    "currency": ["currency"],
    "source_doc": ["sourcedoc", "document"],
}
HISTORY_COLUMNS = {
    "quote_id": ["quoteid"],
    "vendor": ["vendor"],
    "product_id": ["productid", "sku"],
    "quoted_date": ["quoteddate", "date"],
    "quoted_total": ["quotedtotalusd", "total"],
    "status": ["status"],
}

REQUIRED = {
    "po": ["po_id", "line_id", "description", "category", "qty", "metal_g", "set_pcs", "finish"],
    "rate": ["rate_id", "vendor", "effective_from", "metal", "metal_usd_per_gram", "setting_usd_per_stone",
             "labor_usd_per_unit", "handling_rate", "finish_rhodium_usd_per_unit"],
    "quote": ["quote_id"],
    "history": ["quote_id"],
}


def _map_headers(headers: list[Any], spec: dict[str, list[str]]) -> dict[str, int]:
    normed = [norm_header(h) for h in headers]
    out: dict[str, int] = {}
    for canon, aliases in spec.items():
        for cand in [norm_header(canon), *aliases]:
            if cand in normed:
                out[canon] = normed.index(cand)
                break
    return out


def _read_xlsx(path: Path, preferred_sheet: str, issues: list[IngestionIssue]) -> tuple[str, list[Any], list[tuple[int, list[Any]]]]:
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        if preferred_sheet in wb.sheetnames:
            ws = wb[preferred_sheet]
        else:
            ws = wb[wb.sheetnames[0]]
            issues.append(IngestionIssue(code="SHEET_NOT_FOUND", severity=Severity.WARNING, source_file=path.name,
                                         message=f"Sheet '{preferred_sheet}' not found; used '{ws.title}'"))
        rows = list(ws.iter_rows(values_only=True))
        title = ws.title
    finally:
        wb.close()
    if not rows:
        return title, [], []
    header = list(rows[0])
    body = [(i, list(r)) for i, r in enumerate(rows[1:], start=2) if not all(is_blank(c) for c in r)]
    return title, header, body


def _read_csv(path: Path) -> tuple[list[Any], list[tuple[int, list[Any]]]]:
    with path.open(newline="", encoding="utf-8-sig") as fh:
        rows = list(csv.reader(fh))
    if not rows:
        return [], []
    return rows[0], [(i, r) for i, r in enumerate(rows[1:], start=2) if any(c.strip() for c in r)]


def _load_table(
    path: Path,
    kind: str,
    spec: dict[str, list[str]],
    builder: Callable[[dict[str, Any], Evidence], T],
    sheet: str | None,
    issues: list[IngestionIssue],
) -> list[T]:
    if not path.exists():
        issues.append(IngestionIssue(code="FILE_MISSING", severity=Severity.ERROR, source_file=path.name,
                                     message=f"Required input not found: {path}"))
        return []
    try:
        if sheet is not None:
            title, header, body = _read_xlsx(path, sheet, issues)
        else:
            title = "csv"
            header, body = _read_csv(path)
    except Exception as e:  # corrupt / unreadable workbook
        issues.append(IngestionIssue(code="FILE_UNREADABLE", severity=Severity.ERROR, source_file=path.name,
                                     message=f"{type(e).__name__}: {e}"))
        return []
    cols = _map_headers(header, spec)
    missing = [c for c in REQUIRED[kind] if c not in cols]
    if missing:
        issues.append(IngestionIssue(code="MISSING_COLUMNS", severity=Severity.ERROR, source_file=path.name,
                                     location=title, message=f"Missing required columns: {missing}"))
        return []
    out: list[T] = []
    for row_no, cells in body:
        rec = {c: (cells[i] if i < len(cells) else None) for c, i in cols.items()}
        quote = ", ".join(f"{k}={'' if v is None else v}" for k, v in rec.items() if k in ("po_id", "line_id", "rate_id", "quote_id", "vendor", "effective_from", "metal", "description"))
        ev = Evidence(source_file=path.name, sheet_or_page=title, row_or_line=f"row {row_no}", quote=quote[:240])
        try:
            out.append(builder(rec, ev))
        except ValueError as e:
            issues.append(IngestionIssue(code="ROW_INVALID", severity=Severity.ERROR, source_file=path.name,
                                         location=f"{title} row {row_no}", message=str(e)))
    return out


def _po(rec: dict[str, Any], ev: Evidence) -> POLine:
    po_id, line_id = to_text(rec["po_id"]), to_text(rec["line_id"])
    if not po_id or not line_id:
        raise ValueError("po_id/line_id missing")
    qty = to_decimal(rec["qty"], "qty")
    if qty <= 0:
        raise ValueError(f"qty must be positive, got {qty}")
    return POLine(
        po_id=po_id, line_id=line_id, product_id=to_text(rec.get("product_id")),
        description=to_text(rec["description"]) or "", category=to_text(rec["category"]) or "",
        qty=qty, metal_g=to_decimal(rec["metal_g"], "metal_g"),
        stone_ct=None if is_blank(rec.get("stone_ct")) else to_decimal(rec.get("stone_ct"), "stone_ct"),
        set_pcs=to_decimal(rec["set_pcs"], "set_pcs"), finish=to_text(rec["finish"]) or "Standard",
        due_date=None if is_blank(rec.get("due_date")) else to_date(rec.get("due_date"), "due_date"),
        evidence=ev,
    )


def _rate(rec: dict[str, Any], ev: Evidence) -> RateCard:
    frm = to_date(rec["effective_from"], "effective_from")
    to = None if is_blank(rec.get("effective_to")) else to_date(rec.get("effective_to"), "effective_to")
    if to and to < frm:
        raise ValueError(f"effective_to {to} before effective_from {frm}")
    metal = (to_text(rec["metal"]) or "").upper().replace(" ", "")
    return RateCard(
        rate_id=to_text(rec["rate_id"]) or "", vendor=to_text(rec["vendor"]) or "",
        effective_from=frm, effective_to=to, category=to_text(rec.get("category")) or "ALL", metal=metal,
        metal_usd_per_gram=to_decimal(rec["metal_usd_per_gram"], "metal_usd_per_gram"),
        setting_usd_per_stone=to_decimal(rec["setting_usd_per_stone"], "setting_usd_per_stone"),
        labor_usd_per_unit=to_decimal(rec["labor_usd_per_unit"], "labor_usd_per_unit"),
        handling_rate=to_decimal(rec["handling_rate"], "handling_rate"),
        finish_rhodium_usd_per_unit=to_decimal(rec["finish_rhodium_usd_per_unit"], "finish_rhodium_usd_per_unit"),
        notes=to_text(rec.get("notes")) or "", evidence=ev,
    )


def _csv_quote(rec: dict[str, Any], ev: Evidence) -> QuoteRecord:
    qid = to_text(rec["quote_id"])
    if not qid:
        raise ValueError("quote_id missing")
    errors: list[str] = []

    def soft(fn, *a):
        try:
            return fn(*a)
        except ValueError as e:
            errors.append(str(e))
            return None

    ev = ev.model_copy(update={"quote": f"{qid}: {to_text(rec.get('description')) or ''} | qty {rec.get('qty')} {rec.get('unit')} | total {rec.get('quoted_total_usd')}"[:240]})
    return QuoteRecord(
        origin="csv", quote_id=qid, vendor=to_text(rec.get("vendor")),
        quote_date=soft(to_date, rec.get("quote_date"), "quote_date"),
        po_id=to_text(rec.get("po_id")), supplier_line=to_text(rec.get("supplier_line")),
        vendor_sku=to_text(rec.get("vendor_sku")), description=to_text(rec.get("description")),
        qty=soft(to_decimal, rec.get("qty"), "qty"), unit=(to_text(rec.get("unit")) or None),
        quoted_total=soft(to_decimal, rec.get("quoted_total_usd"), "quoted_total_usd"),
        currency=(to_text(rec.get("currency")) or "USD").upper(),
        declared_source_doc=to_text(rec.get("source_doc")), parse_errors=errors, evidence=ev,
    )


def _history(rec: dict[str, Any], ev: Evidence) -> HistoryRecord:
    return HistoryRecord(
        quote_id=to_text(rec["quote_id"]) or "", vendor=to_text(rec.get("vendor")),
        product_id=to_text(rec.get("product_id")),
        quoted_date=None if is_blank(rec.get("quoted_date")) else to_date(rec.get("quoted_date"), "quoted_date"),
        quoted_total=None if is_blank(rec.get("quoted_total")) else to_decimal(rec.get("quoted_total"), "quoted_total"),
        status=to_text(rec.get("status")), evidence=ev,
    )


def load_po_lines(path: Path, issues: list[IngestionIssue]) -> list[POLine]:
    lines = _load_table(path, "po", PO_COLUMNS, _po, "PO_Lines", issues)
    seen: dict[str, POLine] = {}
    for ln in lines:
        if ln.key in seen:
            issues.append(IngestionIssue(code="DUPLICATE_PO_LINE", severity=Severity.ERROR, source_file=path.name,
                                         location=ln.evidence.row_or_line, message=f"{ln.key} duplicated; first occurrence kept"))
            continue
        seen[ln.key] = ln
    return list(seen.values())


def load_rate_cards(path: Path, issues: list[IngestionIssue]) -> list[RateCard]:
    return _load_table(path, "rate", RATE_COLUMNS, _rate, "Rates", issues)


def load_quote_register(path: Path, issues: list[IngestionIssue]) -> list[QuoteRecord]:
    recs = _load_table(path, "quote", QUOTE_COLUMNS, _csv_quote, None, issues)
    for r in recs:
        for err in r.parse_errors:
            issues.append(IngestionIssue(code="FIELD_INVALID", severity=Severity.ERROR, source_file=path.name,
                                         location=r.evidence.row_or_line, message=f"{r.quote_id}: {err}"))
    return recs


def load_history(path: Path, issues: list[IngestionIssue]) -> list[HistoryRecord]:
    return _load_table(path, "history", HISTORY_COLUMNS, _history, "History", issues)
