"""Synthetic package generator for batch / evaluation runs (>=100 quotations with known answers).

Writes the same file layout as the supplied Data/ folder plus gold.json. Gold expected totals are
computed by an INDEPENDENT reference implementation (not audit_agent.calc) so the evaluation
checks the engine against the written policy, not against itself.

Usage:  python eval/synth.py --n 120 --out runs/synth --seed 7
"""
from __future__ import annotations

import argparse
import csv
import json
import random
from datetime import date, timedelta
from decimal import ROUND_HALF_UP, Decimal
from pathlib import Path

from fpdf import FPDF
from openpyxl import Workbook

VENDORS = {"Aurora Craftworks": "A", "Blue Peak Metals": "B", "Crestline Gems": "C", "Dunmore Atelier": "D"}
PRODUCTS = [  # category, sku prefix, PO wording, supplier shorthand wording
    ("Ring", "RNG", "{g} {c} gold diamond solitaire ring", "{g} {cs} RD diamond solitaire ring"),
    ("Ring", "RNG", "{g} {c} gold pavé diamond band", "{g} {c} gold pave diamond band"),
    ("Earrings", "EAR", "{g} {c} gold round diamond stud earrings, one pair", "{g} {cs} RD diamond ear studs / pair"),
    ("Earrings", "EAR", "{g} {c} gold diamond huggie hoops, one pair", "Pair of {g} {cs} round diamond huggies"),
    ("Bracelet", "BRC", "{g} {c} gold diamond tennis bracelet", "{g} {c} gold tennis bracelet w/ diamonds"),
    ("Bracelet", "BRC", "{g} {c} gold sapphire bracelet", "{g} {cs} sapphire bracelet"),
    ("Pendant", "PND", "{g} {c} gold ruby pendant", "{g} {c} gold ruby pendant"),
    ("Pendant", "PND", "{g} {c} gold oval emerald pendant", "{g} {cs} oval emerald pendant"),
    ("Necklace", "NCK", "{g} {c} gold diamond necklace", "{g} {c} gold diamond necklace"),
]
COLOR_SHORT = {"white": "WG", "yellow": "YG", "rose": "RG"}
SCENARIOS = ["exact", "exact", "exact", "overquote", "underquote", "rhodium_extra", "stale_rate", "no_sku", "no_sku_shorthand",
             "unit_piece", "out_of_period", "register_conflict"]
D2 = Decimal("0.01")


def r2(x: Decimal) -> Decimal:
    return x.quantize(D2, rounding=ROUND_HALF_UP)


def ref_total(po: dict, rate: dict) -> Decimal:
    """Reference implementation transcribed from Procurement_Policy_2026 Rules 2-7."""
    metal = r2(Decimal(po["metal_g"]) * Decimal(rate["metal"]))
    setting = r2(Decimal(po["set_pcs"]) * Decimal(rate["setting"]))
    labor = r2(Decimal(rate["labor"]))
    sub = metal + setting + labor
    handling = r2(sub * Decimal(rate["handling"]))
    finish = r2(Decimal(rate["rhodium"])) if po["finish"] == "Rhodium" else Decimal("0")
    return r2(r2(sub + handling + finish) * Decimal(po["qty"]))


def make_rates(rng: random.Random) -> list[dict]:
    rates = []
    for vendor, code in VENDORS.items():
        base14 = rng.choice([38, 40, 41, 42, 43])
        for grade, mult in (("14K", 1), ("18K", Decimal("1.4"))):
            common = {"vendor": vendor, "grade": grade, "rhodium": str(rng.choice([12, 14, 15, 16]))}
            old = {**common, "id": f"{code}-OLD{grade[:2]}", "from": "2026-01-01", "to": "2026-09-30",
                   "metal": str(r2(Decimal(base14) * mult)), "setting": str(rng.choice(["2.5", "2.75", "3"])),
                   "labor": str(rng.choice([35, 38, 40]) + (10 if grade == "18K" else 0)), "handling": rng.choice(["0.07", "0.08", "0.09"])}
            new = {**common, "id": f"{code}-NEW{grade[:2]}", "from": "2026-10-01", "to": "2026-12-31",
                   "metal": str(r2(Decimal(old["metal"]) + rng.choice([2, 3, 4]))), "setting": str(Decimal(old["setting"]) + Decimal("0.5")),
                   "labor": str(int(old["labor"]) + 5), "handling": "0.1"}
            rates += [old, new]
    return rates


def pick_rate(rates, vendor, grade, d):
    for r in rates:
        if r["vendor"] == vendor and r["grade"] == grade and r["from"] <= d <= r["to"]:
            return r
    return None


def generate(n: int, out: Path, seed: int = 7) -> dict:
    rng = random.Random(seed)
    rates = make_rates(rng)
    pos, quotes, gold = [], [], {}
    for i in range(1, n + 1):
        cat, prefix, po_tpl, q_tpl = rng.choice(PRODUCTS)
        grade, color = rng.choice(["14K", "18K"]), rng.choice(["white", "yellow", "rose"])
        vendor = rng.choice(list(VENDORS))
        po_id, line = f"PO-{3000 + (i - 1) // 10}", f"L{(i - 1) % 10 + 1:02d}"
        qd = date(2026, 9, 1) + timedelta(days=rng.randint(0, 110))
        po = {"po_id": po_id, "line_id": line, "product_id": f"{prefix}-{1000 + i}", "category": cat,
              "description": po_tpl.format(g=grade, c=color), "qty": rng.randint(1, 4),
              "metal_g": str(Decimal(rng.randint(25, 140)) / 10), "stone_ct": "1", "set_pcs": rng.choice([1, 2, 12, 18, 24, 35]),
              "finish": "Rhodium" if (color == "white" and rng.random() < 0.3) else "Standard", "due_date": str(qd + timedelta(days=7))}
        pos.append(po)
        scen = rng.choice(SCENARIOS)
        if scen == "out_of_period":
            qd = date(2027, 1, rng.randint(2, 28))
        if scen == "stale_rate":
            qd = date(2026, 10, rng.randint(1, 31))
        rate = pick_rate(rates, vendor, grade, str(qd))
        q = {"quote_id": f"S-{i:04d}", "vendor": vendor, "quote_date": str(qd), "po_id": po_id, "supplier_line": f"S-{(i - 1) % 10 + 1:02d}",
             "vendor_sku": po["product_id"], "description": po["description"], "qty": po["qty"],
             "unit": "pair" if cat == "Earrings" else "piece"}
        expected = ref_total(po, rate) if rate else None
        status, quoted = "WITHIN_TOLERANCE", expected
        if scen in ("no_sku", "no_sku_shorthand"):
            q["vendor_sku"] = ""
            if scen == "no_sku_shorthand":
                q["description"] = q_tpl.format(g=grade.lower() if rng.random() < 0.5 else grade, c=color, cs=COLOR_SHORT[color])
        if scen == "overquote":
            quoted, status = expected + Decimal(rng.choice(["5", "12.50", "40", "75"])), "FLAG"
        elif scen == "underquote":
            quoted, status = expected - Decimal(rng.choice(["1", "9.99", "20"])), "FLAG"
        elif scen == "rhodium_extra" and po["finish"] == "Standard":
            quoted, status = r2(expected + Decimal(rate["rhodium"]) * po["qty"]), "FLAG"
        elif scen == "stale_rate":
            old = next(r for r in rates if r["vendor"] == vendor and r["grade"] == grade and r["to"] == "2026-09-30")
            quoted = ref_total(po, old)
            status = "FLAG" if abs(quoted - expected) > D2 else "WITHIN_TOLERANCE"
        elif scen == "unit_piece" and cat == "Earrings":
            q["unit"], status = "piece", "NEEDS_REVIEW"
        elif scen == "out_of_period":
            quoted, status = Decimal(rng.randint(200, 900)) + Decimal("0.50"), "NEEDS_REVIEW"
        q["quoted_total_usd"] = str(quoted)
        q["pdf_total"] = str(quoted)
        if scen == "register_conflict":
            q["quoted_total_usd"], status = str(quoted + Decimal("10.00")), "NEEDS_REVIEW"
        quotes.append(q)
        gold[q["quote_id"]] = {"po_line": f"{po_id}/{line}", "rate_id": rate["id"] if rate else None,
                               "expected": str(expected) if (expected is not None and status != "NEEDS_REVIEW") else None,
                               "quoted": str(quoted), "status": status, "scenario": scen,
                               "fields": {k: str(q[k]) for k in ("quote_date", "po_id", "vendor_sku", "qty", "unit")} | {"total": str(quoted)}}
    _write(out, pos, rates, quotes)
    (out / "gold.json").write_text(json.dumps(gold, indent=1))
    return gold


def _write(out: Path, pos, rates, quotes):
    (out / "Quotations").mkdir(parents=True, exist_ok=True)
    (out / "Contracts").mkdir(parents=True, exist_ok=True)
    wb = Workbook(); ws = wb.active; ws.title = "PO_Lines"
    cols = ["po_id", "line_id", "product_id", "description", "category", "qty", "metal_g", "stone_ct", "set_pcs", "finish", "due_date"]
    ws.append(cols)
    for p in pos:
        ws.append([float(p[c]) if c in ("metal_g", "stone_ct") else p[c] for c in cols])
    wb.save(out / "Purchase_Orders.xlsx")
    wb = Workbook(); ws = wb.active; ws.title = "Rates"
    ws.append(["RateID", "Vendor", "EffectiveFrom", "EffectiveTo", "Category", "Metal", "MetalUSDperGram", "SettingUSDperStone",
               "LaborUSDperUnit", "HandlingRate", "FinishRhodiumUSDperUnit", "Notes"])
    for r in rates:
        ws.append([r["id"], r["vendor"], r["from"], r["to"], "ALL", r["grade"], float(r["metal"]), float(r["setting"]),
                   float(r["labor"]), float(r["handling"]), float(r["rhodium"]), "synthetic"])
    wb.save(out / "Vendor_Rate_Cards.xlsx")
    with (out / "Quote_Register.csv").open("w", newline="") as fh:
        w = csv.writer(fh)
        w.writerow(["quote_id", "vendor", "quote_date", "po_id", "supplier_line", "vendor_sku", "description", "qty", "unit",
                    "quoted_total_usd", "currency", "source_doc"])
        for q in quotes:
            w.writerow([q["quote_id"], q["vendor"], q["quote_date"], q["po_id"], q["supplier_line"], q["vendor_sku"], q["description"],
                        q["qty"], q["unit"], q["quoted_total_usd"], "USD", f"Synth_{VENDORS[q['vendor']]}.pdf"])
    for vendor in VENDORS:
        pdf = FPDF(); pdf.set_auto_page_break(True, 15); pdf.add_page(); pdf.set_font("Helvetica", size=10)
        pdf.cell(0, 6, f"{vendor} | COMMERCIAL QUOTATIONS", new_x="LMARGIN", new_y="NEXT")
        for q in (q for q in quotes if q["vendor"] == vendor):
            total = f"{Decimal(q['pdf_total']):,.2f}"
            for ln in (f"{q['quote_id']} Date: {q['quote_date']} PO: {q['po_id']} Line: {q['supplier_line']}",
                       f"Item: {q['description'].replace('é', 'e')}",
                       f"SKU: {q['vendor_sku'] or '[not provided]'} QTY: {q['qty']} {q['unit']} Total: USD {total}",
                       "Payment: Net 30 days."):
                pdf.cell(0, 5, ln, new_x="LMARGIN", new_y="NEXT")
        pdf.output(str(out / "Quotations" / f"Synth_{vendor.replace(' ', '_')}.pdf"))
    import shutil
    src = Path(__file__).resolve().parents[1] / "Data" / "Contracts" / "Procurement_Policy_2026.pdf"
    shutil.copy(src, out / "Contracts" / src.name)


if __name__ == "__main__":
    ap = argparse.ArgumentParser()
    ap.add_argument("--n", type=int, default=120)
    ap.add_argument("--out", type=Path, default=Path("runs/synth"))
    ap.add_argument("--seed", type=int, default=7)
    a = ap.parse_args()
    g = generate(a.n, a.out, a.seed)
    print(f"wrote {len(g)} synthetic quotations to {a.out}")
